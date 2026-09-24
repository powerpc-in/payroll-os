"""Reporting engine — reusable dataset definitions + filters/grouping/aggregation +
CSV/Excel/PDF export, plus saved (custom) reports. No hard-coded report strings
in the UI: the frontend renders whatever this engine defines."""

import csv
import io
from copy import deepcopy
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from lib.auth import Context, new_id, require_perm
from lib.db import db

router = APIRouter(prefix="/v1/reports", tags=["Reports"])


def _f(key: str, label: str, numeric: bool = False) -> dict:
    return {"key": key, "label": label, "numeric": numeric}


def _pay_fields(extra: list[dict] | None = None) -> list[dict]:
    base = [
        _f("period", "Period"), _f("employee_code", "Employee ID"), _f("employee_name", "Employee"),
        _f("department", "Department"), _f("location", "Location"),
        _f("gross_earnings", "Gross", True), _f("total_deductions", "Deductions", True),
    ]
    return base + (extra or [])


DATASETS: dict[str, dict] = {
    "payroll_register": {"name": "Payroll Register", "collection": "payroll_employees", "time_field": "period",
                         "fields": _pay_fields([_f("net_pay", "Net Pay", True), _f("employer_cost", "Employer Cost", True)])},
    "tds": {"name": "TDS", "collection": "payroll_employees", "time_field": "period",
            "fields": [_f("period", "Period"), _f("employee_code", "Employee ID"), _f("employee_name", "Employee"),
                       _f("taxable_gross", "Taxable Gross", True), _f("tds_amount", "TDS", True), _f("fy", "Financial Year")]},
    "pf": {"name": "Provident Fund", "collection": "payroll_employees", "time_field": "period",
           "fields": [_f("period", "Period"), _f("employee_code", "Employee ID"), _f("employee_name", "Employee"),
                      _f("pf_employee", "Employee PF", True), _f("pf_employer", "Employer PF", True)]},
    "esi": {"name": "ESI", "collection": "payroll_employees", "time_field": "period",
            "fields": [_f("period", "Period"), _f("employee_code", "Employee ID"), _f("employee_name", "Employee"),
                       _f("esi_employee", "Employee ESI", True), _f("esi_employer", "Employer ESI", True)]},
    "pt": {"name": "Professional Tax", "collection": "payroll_employees", "time_field": "period",
           "fields": [_f("period", "Period"), _f("employee_code", "Employee ID"), _f("employee_name", "Employee"),
                      _f("state", "State"), _f("pt_amount", "PT", True)]},
    "lwf": {"name": "Labour Welfare Fund", "collection": "payroll_employees", "time_field": "period",
            "fields": [_f("period", "Period"), _f("employee_code", "Employee ID"), _f("employee_name", "Employee"),
                       _f("state", "State"), _f("lwf_amount", "LWF", True)]},
    "employer_contributions": {"name": "Employer Contributions", "collection": "payroll_employees",
                               "time_field": "period",
                               "fields": [_f("period", "Period"), _f("employee_code", "Employee ID"),
                                          _f("employee_name", "Employee"), _f("pf_employer", "Employer PF", True),
                                          _f("esi_employer", "Employer ESI", True),
                                          _f("total_employer_contributions", "Total Contributions", True),
                                          _f("employer_cost", "Employer Cost", True)]},
    "net_pay": {"name": "Net Pay", "collection": "payroll_employees", "time_field": "period",
                "fields": _pay_fields([_f("tds_amount", "TDS", True), _f("net_pay", "Net Pay", True)])},
    "employee_earnings": {"name": "Employee Earnings (by component)", "collection": "payroll_employees",
                          "time_field": "period", "flatten": ("earnings",),
                          "fields": [_f("period", "Period"), _f("employee_code", "Employee ID"),
                                     _f("employee_name", "Employee"), _f("component", "Component"),
                                     _f("amount", "Amount", True)]},
    "employee_deductions": {"name": "Employee Deductions (by component)", "collection": "payroll_employees",
                            "time_field": "period", "flatten": ("deductions",),
                            "fields": [_f("period", "Period"), _f("employee_code", "Employee ID"),
                                       _f("employee_name", "Employee"), _f("component", "Component"),
                                       _f("amount", "Amount", True)]},
    "ctc": {"name": "CTC (annual, from assignments)", "collection": "salary_assignments", "time_field": None,
            "fields": [_f("employee_code", "Employee ID"), _f("employee_name", "Employee"),
                       _f("department", "Department"), _f("gross_monthly", "Gross / month", True),
                       _f("annual_ctc", "Annual CTC", True)]},
    "department_payroll": {"name": "Department Payroll", "collection": "payroll_employees",
                           "time_field": "period", "group_only": "department",
                           "fields": [_f("group", "Department"), _f("headcount", "Employees", True),
                                      _f("gross_earnings", "Gross", True), _f("net_pay", "Net Pay", True),
                                      _f("employer_cost", "Employer Cost", True)]},
    "location_payroll": {"name": "Location Payroll", "collection": "payroll_employees",
                         "time_field": "period", "group_only": "location",
                         "fields": [_f("group", "Location"), _f("headcount", "Employees", True),
                                    _f("gross_earnings", "Gross", True), _f("net_pay", "Net Pay", True),
                                    _f("employer_cost", "Employer Cost", True)]},
    "cost_centre_payroll": {"name": "Cost Centre Payroll", "collection": "payroll_employees",
                            "time_field": "period", "group_only": "cost_centre",
                            "fields": [_f("group", "Cost Centre"), _f("headcount", "Employees", True),
                                       _f("gross_earnings", "Gross", True), _f("net_pay", "Net Pay", True),
                                       _f("employer_cost", "Employer Cost", True)]},
    "reimbursements": {"name": "Reimbursements", "collection": "reimbursements", "time_field": None,
                       "fields": [_f("date", "Date"), _f("employee_name", "Employee"),
                                  _f("category", "Category"), _f("amount", "Claimed", True),
                                  _f("approved_amount", "Approved", True), _f("status", "Status")]},
    "loans": {"name": "Loans & Advances", "collection": "loans", "time_field": None,
              "fields": [_f("employee_name", "Employee"), _f("name", "Loan"), _f("principal", "Principal", True),
                         _f("emi", "EMI", True), _f("outstanding", "Outstanding", True),
                         _f("interest_rate", "Interest %", True), _f("status", "Status")]},
    "ff": {"name": "Full & Final Settlements", "collection": "ffs", "time_field": None,
           "fields": [_f("employee_name", "Employee"), _f("employee_code", "Employee ID"),
                      _f("department", "Department"), _f("last_working_day", "LWD"),
                      _f("total_payable", "Payable", True), _f("total_recoveries", "Recoveries", True),
                      _f("tax_adjustment", "Tax adjustment", True),
                      _f("net_settlement", "Net Settlement", True), _f("status", "Status")]},
    "attendance": {"name": "Attendance summary", "collection": "attendance", "time_field": "date",
                   "fields": [_f("employee_code", "Employee ID"), _f("employee_name", "Employee"),
                              _f("present", "Present", True), _f("absent", "Absent", True),
                              _f("half_day", "Half days", True), _f("paid_leave", "Paid leave", True),
                              _f("unpaid_leave", "Unpaid leave", True),
                              _f("overtime_hours", "OT hours", True), _f("lop", "LOP days", True)]},
    "leave": {"name": "Leave requests", "collection": "leave_requests", "time_field": None,
              "fields": [_f("employee_name", "Employee"), _f("leave_type_name", "Leave type"),
                         _f("from_date", "From"), _f("to_date", "To"),
                         _f("days", "Days", True), _f("paid", "Paid"), _f("status", "Status")]},
}

# These values reveal employee pay, statutory deductions, or employer payroll cost.
COMPENSATION_FIELDS = {
    "gross_earnings", "total_deductions", "net_pay", "employer_cost", "taxable_gross",
    "tds_amount", "pf_employee", "pf_employer", "esi_employee", "esi_employer",
    "pt_amount", "lwf_amount", "total_employer_contributions", "amount",
    "gross_monthly", "annual_ctc",
}


def _visible_dataset(ctx: Context, dataset: str) -> dict:
    cfg = DATASETS.get(dataset)
    if not cfg:
        raise HTTPException(status_code=404, detail=f"Unknown dataset '{dataset}'")
    visible = deepcopy(cfg)
    if not ctx.can("salary.view"):
        visible["fields"] = [f for f in visible["fields"] if f["key"] not in COMPENSATION_FIELDS]
    return visible


def _validate_report_fields(r: "ReportRun", cfg: dict) -> None:
    available = {f["key"] for f in cfg["fields"]}
    requested = set(r.fields or [])
    hidden = requested - available
    hidden.update(f.field for f in r.filters if f.field not in available)
    if r.group_by and r.group_by not in available and r.group_by != cfg.get("group_only"):
        hidden.add(r.group_by)
    if r.sort_by and r.sort_by not in available and r.sort_by not in ("group", "headcount"):
        hidden.add(r.sort_by)
    if hidden:
        raise HTTPException(status_code=403, detail="Report field is not available to this role")

AGGREGATIONS = ["sum", "avg", "min", "max", "count"]
FILTER_OPS = ["eq", "ne", "contains", "gt", "gte", "lt", "lte"]


class ReportFilter(BaseModel):
    field: str
    op: str = "eq"
    value: str = ""


class ReportRun(BaseModel):
    dataset: str
    fields: list[str] | None = None          # column subset; None → every dataset field
    filters: list[ReportFilter] = []         # generic field/op/value filters
    period_from: str | None = None
    period_to: str | None = None
    department: str | None = None
    location: str | None = None
    group_by: str | None = None
    aggregate: str = "sum"                   # applied to numeric columns when grouping
    sort_by: str | None = None
    sort_dir: str = "asc"
    limit: int = 500


async def _rows_for(ctx: Context, cfg: dict, r: ReportRun) -> list[dict]:
    collection = db[cfg["collection"]]
    query: dict = {"org_id": ctx.org_id}
    time_field = cfg.get("time_field")
    if time_field:
        rng: dict = {}
        if r.period_from:
            rng["$gte"] = r.period_from
        if r.period_to:
            rng["$lte"] = r.period_to
        if rng:
            query[time_field] = rng
    docs = await collection.find(query).to_list(r.limit * 4 + 1000)

    flat = cfg.get("flatten")
    rows: list[dict] = []
    if flat:
        for d in docs:
            snap = dict(d.get("employee_snapshot") or {})
            base = {"period": d.get("period"), "employee_code": snap.get("employee_code"),
                    "employee_name": snap.get("name"), "department": snap.get("department"),
                    "location": snap.get("location")}
            for item in d.get(flat[0], []):
                row = dict(base)
                row["component"] = f"{item.get('code', '')} · {item.get('name', '')}"
                row["amount"] = item.get("amount", 0)
                rows.append(row)
    elif cfg["collection"] == "payroll_employees":
        for d in docs:
            snap = dict(d.pop("employee_snapshot", {}) or {})
            d.update(snap)
            d["employee_name"] = snap.get("name")
            rows.append(d)
    elif cfg["collection"] == "salary_assignments":
        active = [d for d in docs if d.get("active", True)]
        emps = {e["id"]: e for e in await db.employees.find(
            {"org_id": ctx.org_id, "id": {"$in": [a["employee_id"] for a in active]}}).to_list(5000)}
        for a in active:
            e = emps.get(a["employee_id"], {})
            rows.append({
                "employee_code": e.get("employee_code"), "employee_name": e.get("name"),
                "department": e.get("department_name"), "gross_monthly": a.get("gross_monthly"),
                "annual_ctc": round((a.get("gross_monthly") or 0) * 12, 2),
            })
    elif cfg["collection"] == "attendance":
        summary: dict[str, dict] = {}
        for d in docs:
            entry = summary.setdefault(d["employee_id"], {
                "present": 0.0, "absent": 0.0, "half_day": 0.0, "paid_leave": 0.0,
                "unpaid_leave": 0.0, "overtime_hours": 0.0})
            st = d.get("status")
            if st in entry and st != "overtime_hours":
                entry[st] += d.get("days", 1)
            entry["overtime_hours"] += d.get("overtime_hours", 0) or 0
        emps = {e["id"]: e for e in await db.employees.find(
            {"org_id": ctx.org_id, "id": {"$in": list(summary.keys())}}).to_list(5000)}
        for emp_id, e in summary.items():
            emp = emps.get(emp_id, {})
            rows.append({
                "employee_code": emp.get("employee_code"), "employee_name": emp.get("name"),
                **e, "lop": e["absent"] + e["unpaid_leave"] + 0.5 * e["half_day"],
            })
        if not r.sort_by:
            rows.sort(key=lambda x: x.get("employee_name") or "")
    else:
        rows = docs

    if r.department:
        rows = [x for x in rows if x.get("department") == r.department]
    if r.location:
        rows = [x for x in rows if x.get("location") == r.location]
    for flt in r.filters:
        rows = [x for x in rows if _matches(x.get(flt.field), flt)]
    return rows


def _matches(value, flt: ReportFilter) -> bool:
    """Generic field/op/value predicate. Numeric ops coerce; text ops compare lowercased."""
    target = flt.value
    if flt.op in ("gt", "gte", "lt", "lte"):
        try:
            left, right = float(value), float(target)
        except (TypeError, ValueError):
            return False
        return {"gt": left > right, "gte": left >= right,
                "lt": left < right, "lte": left <= right}[flt.op]
    left_s = "" if value is None else str(value).lower()
    right_s = str(target).lower()
    if flt.op == "contains":
        return right_s in left_s
    if flt.op == "ne":
        return left_s != right_s
    return left_s == right_s


def _aggregate(values: list[float], how: str) -> float:
    if not values:
        return 0.0
    if how == "avg":
        return round(sum(values) / len(values), 2)
    if how == "min":
        return round(min(values), 2)
    if how == "max":
        return round(max(values), 2)
    if how == "count":
        return float(len(values))
    return round(sum(values), 2)


def _project(rows: list[dict], cfg: dict, r: ReportRun) -> tuple[list[str], list[dict], dict]:
    fields = cfg["fields"]
    if r.fields:
        chosen = [f for f in fields if f["key"] in r.fields]
        if chosen:
            fields = chosen
    how = r.aggregate if r.aggregate in AGGREGATIONS else "sum"
    if cfg.get("group_only"):
        r.group_by = cfg["group_only"]
    if r.group_by:
        buckets: dict[str, dict[str, list[float]]] = {}
        counts: dict[str, int] = {}
        for row in rows:
            key = str(row.get(r.group_by) or "—")
            counts[key] = counts.get(key, 0) + 1
            bucket = buckets.setdefault(key, {})
            for f in fields:
                if f["numeric"] and f["key"] != "headcount":
                    bucket.setdefault(f["key"], []).append(float(row.get(f["key"]) or 0))
        out = []
        for key, bucket in buckets.items():
            grouped_row: dict = {"group": key, "headcount": counts[key]}
            for f in fields:
                if f["numeric"] and f["key"] != "headcount":
                    grouped_row[f["key"]] = _aggregate(bucket.get(f["key"], []), how)
            out.append(grouped_row)
        columns = ["group", "headcount"] + [f["key"] for f in fields
                                            if f["numeric"] and f["key"] not in ("group", "headcount")]
        if r.sort_by in columns:
            out.sort(key=lambda x: (x.get(r.sort_by) is None, x.get(r.sort_by)),
                     reverse=r.sort_dir == "desc")
        else:
            out.sort(key=lambda x: x["group"])
        totals = {c: round(sum(float(x.get(c) or 0) for x in out), 2)
                  for c in columns if c != "group"}
        return columns, out, totals

    keys = [f["key"] for f in fields]
    out = [{k: row.get(k) for k in keys} for row in rows]
    if r.sort_by:
        out.sort(key=lambda x: (x.get(r.sort_by) is None, x.get(r.sort_by)),
                 reverse=r.sort_dir == "desc")
    totals = {f["key"]: round(sum(x.get(f["key"]) or 0 for x in out if isinstance(x.get(f["key"]), (int, float))), 2)
              for f in fields if f["numeric"]}
    return keys, out[: r.limit], totals


def _export_response(columns: list[str], rows: list[dict], fmt: str, name: str):
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(columns)
        for row in rows:
            writer.writerow([row.get(c, "") for c in columns])
        return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                                 headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})
    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(columns)
        for row in rows:
            ws.append([row.get(c) for c in columns])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return StreamingResponse(iter([bio.getvalue()]),
                                 media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})
    if fmt == "pdf":
        from fpdf import FPDF
        pdf = FPDF(orientation="L", format="A4")
        pdf.add_page()
        pdf.set_font("Helvetica", size=8)
        pdf.cell(0, 6, name, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=6.5)
        pdf.cell(0, 4, datetime.now(timezone.utc).strftime("Generated %Y-%m-%d %H:%M UTC"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
        col_w = 270 / max(1, len(columns))
        pdf.set_fill_color(241, 245, 249)
        for c in columns:
            pdf.cell(col_w, 5, str(c)[:28], border=1, fill=True)
        pdf.ln()
        for row in rows[:400]:
            for c in columns:
                v = row.get(c, "")
                if isinstance(v, float):
                    text = f"{v:,.2f}"
                else:
                    text = str(v if v is not None else "")[:28]
                pdf.cell(col_w, 5, text, border=1)
            pdf.ln()
        return StreamingResponse(iter([bytes(pdf.output())]), media_type="application/pdf",
                                 headers={"Content-Disposition": f'attachment; filename="{name}.pdf"'})
    raise HTTPException(status_code=422, detail="format must be csv, xlsx or pdf")


GROUPABLE = ["department", "location", "cost_centre", "employee_code", "period", "state", "status"]


@router.get("/datasets")
async def datasets(ctx: Context = Depends(require_perm("reports.view"))):
    visible_datasets = {key: _visible_dataset(ctx, key) for key in DATASETS}
    return {
        "datasets": [{"key": k, "name": v["name"],
                      "fields": v["fields"], "time_field": v.get("time_field"),
                      "group_only": v.get("group_only"), "flatten": bool(v.get("flatten")),
                      "groupable": [g for g in GROUPABLE
                                    if g in {f["key"] for f in v["fields"]} or g == v.get("group_only")]}
                     for k, v in visible_datasets.items()],
        "aggregations": AGGREGATIONS,
        "filter_ops": FILTER_OPS,
    }


@router.post("/run")
async def run_report(r: ReportRun, ctx: Context = Depends(require_perm("reports.view"))):
    cfg = _visible_dataset(ctx, r.dataset)
    _validate_report_fields(r, cfg)
    rows = await _rows_for(ctx, cfg, r)
    columns, out, totals = _project(rows, cfg, r)
    labels = {f["key"]: f["label"] for f in cfg["fields"]}
    labels.update({"group": (r.group_by or "Group").replace("_", " ").title(),
                   "headcount": "Rows"})
    return {"dataset": r.dataset, "name": cfg["name"], "columns": columns,
            "labels": labels, "numeric": [f["key"] for f in cfg["fields"] if f["numeric"]]
            + ["headcount"],
            "group_by": r.group_by, "aggregate": r.aggregate,
            "rows": out, "totals": totals, "count": len(out)}


@router.post("/export")
async def export_report(r: ReportRun, format: str = "csv",
                        ctx: Context = Depends(require_perm("reports.export"))):
    cfg = _visible_dataset(ctx, r.dataset)
    _validate_report_fields(r, cfg)
    rows = await _rows_for(ctx, cfg, r)
    columns, out, _ = _project(rows, cfg, r)
    return _export_response(columns, out, format, cfg["name"])


class SavedReportIn(BaseModel):
    name: str
    config: dict
    visualization: str = "table"   # table | bar | line — chosen by the builder UI


@router.get("/saved")
async def list_saved(ctx: Context = Depends(require_perm("reports.view"))):
    docs = await db.saved_reports.find({"org_id": ctx.org_id}).sort("created_at", -1).to_list(100)
    for d in docs:
        d.pop("_id", None)
    return docs


@router.post("/saved")
async def save_report(input: SavedReportIn, ctx: Context = Depends(require_perm("reports.manage"))):
    doc = {"id": new_id(), "org_id": ctx.org_id, "name": input.name, "config": input.config,
           "visualization": input.visualization,
           "created_by": ctx.user["email"], "created_at": datetime.now(timezone.utc)}
    await db.saved_reports.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.put("/saved/{saved_id}")
async def update_saved(saved_id: str, input: SavedReportIn,
                       ctx: Context = Depends(require_perm("reports.manage"))):
    res = await db.saved_reports.update_one(
        {"id": saved_id, "org_id": ctx.org_id},
        {"$set": {"name": input.name, "config": input.config,
                  "visualization": input.visualization,
                  "updated_at": datetime.now(timezone.utc)}})
    if not res.matched_count:
        raise HTTPException(status_code=404, detail="Saved report not found")
    return await db.saved_reports.find_one({"id": saved_id, "org_id": ctx.org_id}, {"_id": 0})


@router.post("/saved/{saved_id}/run")
async def run_saved(saved_id: str, ctx: Context = Depends(require_perm("reports.view"))):
    """Runs a stored configuration against live data — saved reports are configurations,
    never cached result sets."""
    doc = await db.saved_reports.find_one({"id": saved_id, "org_id": ctx.org_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Saved report not found")
    try:
        r = ReportRun(**{k: v for k, v in (doc.get("config") or {}).items()
                         if k in ReportRun.model_fields})
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Saved configuration is invalid: {exc}") from exc
    return await run_report(r, ctx)


@router.delete("/saved/{saved_id}")
async def delete_saved(saved_id: str, ctx: Context = Depends(require_perm("reports.manage"))):
    await db.saved_reports.delete_one({"id": saved_id, "org_id": ctx.org_id})
    return {"ok": True}
